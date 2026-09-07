#!/usr/bin/env python3
"""
Fusionne les corrections de taggage PASA (colonne Sous-Action) dans l'export ETPT complet.

Usage :
  python3 scripts/merge_pasa_corrections.py
  python3 scripts/merge_pasa_corrections.py trdata/base.xlsx trdata/corrections.xlsx trdata/output.xlsx
"""

from __future__ import annotations

import sys
from pathlib import Path
from typing import Any, Optional

import openpyxl

SHEET_BASE = "Données annuelles"
SHEET_CORR = "Feuil1"
HEADER_ROW = 5

# Indices dans le fichier corrections (sans en-tête, 0-based)
CORR = {
    "nir": 0,
    "matricule": 1,
    "nom": 2,
    "nom_usage": 3,
    "prenom": 4,
    "annee": 5,
    "mois": 6,
    "date_affect": 7,
    "niveau03": 8,
    "niveau06": 9,
    "niveau08": 10,
    "uo_long": 11,
    "statut": 12,
    "qual_stat": 13,
    "position_date": 14,
    "position_code": 15,
    "position_lib": 16,
    "poste_code": 17,
    "poste": 18,
    "programme": 19,
    "action": 20,
    "sous_action": 21,
    "taux": 22,
    "grade14": 23,
    "grade3": 24,
    "etpt": 25,
    "etpt_prog": 26,
    "nne": 27,
    "grade_nne": 28,
}

# Indices dans le fichier base (0-based, colonnes A-B vides)
BASE = {
    "nir": 2,
    "matricule": 3,
    "nom": 4,
    "nom_usage": 5,
    "prenom": 6,
    "annee": 7,
    "mois": 8,
    "date_affect": 9,
    "niveau03": 10,
    "niveau06": 11,
    "niveau08": 12,
    "uo_long": 13,
    "statut": 14,
    "qual_stat": 15,
    "position_date": 16,
    "position_code": 17,
    "position_lib": 18,
    "poste_code": 19,
    "poste": 20,
    "programme": 21,
    "action": 22,
    "sous_action": 23,
    "taux": 24,
    "grade14": 25,
    "grade3": 26,
    "etpt": 27,
    "etpt_prog": 28,
    "nne": 29,
    "grade_nne": 30,
}


def _txt(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _get(row: tuple[Any, ...], idx: int) -> Any:
    return row[idx] if idx < len(row) else None


def corr_to_base_row(corr: tuple[Any, ...]) -> list[Any]:
    """Convertit une ligne corrections vers le format du fichier base (31 colonnes)."""
    row: list[Any] = [None, None]
    mapping = [
        (CORR["nir"], BASE["nir"]),
        (CORR["matricule"], BASE["matricule"]),
        (CORR["nom"], BASE["nom"]),
        (CORR["nom_usage"], BASE["nom_usage"]),
        (CORR["prenom"], BASE["prenom"]),
        (CORR["annee"], BASE["annee"]),
        (CORR["mois"], BASE["mois"]),
        (CORR["date_affect"], BASE["date_affect"]),
        (CORR["niveau03"], BASE["niveau03"]),
        (CORR["niveau06"], BASE["niveau06"]),
        (CORR["niveau08"], BASE["niveau08"]),
        (CORR["uo_long"], BASE["uo_long"]),
        (CORR["statut"], BASE["statut"]),
        (CORR["qual_stat"], BASE["qual_stat"]),
        (CORR["position_date"], BASE["position_date"]),
        (CORR["position_code"], BASE["position_code"]),
        (CORR["position_lib"], BASE["position_lib"]),
        (CORR["poste_code"], BASE["poste_code"]),
        (CORR["poste"], BASE["poste"]),
        (CORR["programme"], BASE["programme"]),
        (CORR["action"], BASE["action"]),
        (CORR["sous_action"], BASE["sous_action"]),
        (CORR["taux"], BASE["taux"]),
        (CORR["grade14"], BASE["grade14"]),
        (CORR["grade3"], BASE["grade3"]),
        (CORR["etpt"], BASE["etpt"]),
        (CORR["etpt_prog"], BASE["etpt_prog"]),
        (CORR["nne"], BASE["nne"]),
        (CORR["grade_nne"], BASE["grade_nne"]),
    ]
    while len(row) < 31:
        row.append(None)
    for c_idx, b_idx in mapping:
        row[b_idx] = _get(corr, c_idx)
    return row


def find_match(
    corr: tuple[Any, ...],
    index: dict[tuple[str, int], list[int]],
    rows: list[tuple[Any, ...]],
) -> Optional[int]:
    mat = _txt(_get(corr, CORR["matricule"]))
    try:
        mois = int(float(_get(corr, CORR["mois"])))
    except (TypeError, ValueError):
        return None
    poste = _txt(_get(corr, CORR["poste"]))
    poste_code = _txt(_get(corr, CORR["poste_code"]))

    candidates = index.get((mat, mois), [])
    for ri in candidates:
        if _txt(_get(rows[ri], BASE["poste"])) == poste:
            return ri
    for ri in candidates:
        if poste_code and _txt(_get(rows[ri], BASE["poste_code"])) == poste_code:
            return ri
    for ri in candidates:
        old_poste = _txt(_get(rows[ri], BASE["poste"]))
        if poste[:40] and old_poste[:40] == poste[:40]:
            return ri
    if len(candidates) == 1:
        return candidates[0]
    return None


def merge(base_path: Path, corr_path: Path, out_path: Path) -> dict[str, int]:
    wb = openpyxl.load_workbook(base_path)
    if SHEET_BASE not in wb.sheetnames:
        raise SystemExit(f"Feuille « {SHEET_BASE} » absente dans {base_path}")

    ws = wb[SHEET_BASE]
    all_rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    header_rows = all_rows[:HEADER_ROW]
    data_rows = list(all_rows[HEADER_ROW:])

    wb_corr = openpyxl.load_workbook(corr_path, read_only=True, data_only=True)
    corr_sheet = SHEET_CORR if SHEET_CORR in wb_corr.sheetnames else wb_corr.sheetnames[0]
    corrections = list(wb_corr[corr_sheet].iter_rows(values_only=True))

    index: dict[tuple[str, int], list[int]] = {}
    for ri, row in enumerate(data_rows):
        mat = _txt(_get(row, BASE["matricule"]))
        try:
            mois = int(float(_get(row, BASE["mois"])))
        except (TypeError, ValueError):
            continue
        index.setdefault((mat, mois), []).append(ri)

    stats = {"updated": 0, "inserted": 0, "unchanged": 0, "total_corrections": len(corrections)}

    for corr in corrections:
        sa_new = _txt(_get(corr, CORR["sous_action"]))
        act_new = _txt(_get(corr, CORR["action"]))
        target = find_match(corr, index, data_rows)

        if target is not None:
            row = list(data_rows[target])
            old_sa = _txt(_get(tuple(row), BASE["sous_action"]))
            old_act = _txt(_get(tuple(row), BASE["action"]))
            changed = False
            if sa_new and sa_new != old_sa:
                row[BASE["sous_action"]] = sa_new
                changed = True
            if act_new and act_new != old_act:
                row[BASE["action"]] = act_new
                changed = True
            if changed:
                data_rows[target] = tuple(row)
                stats["updated"] += 1
            else:
                stats["unchanged"] += 1
        else:
            new_row = tuple(corr_to_base_row(corr))
            data_rows.append(new_row)
            mat = _txt(_get(corr, CORR["matricule"]))
            try:
                mois = int(float(_get(corr, CORR["mois"])))
            except (TypeError, ValueError):
                mois = 0
            index.setdefault((mat, mois), []).append(len(data_rows) - 1)
            stats["inserted"] += 1

    # Réécrire la feuille
    ws.delete_rows(1, ws.max_row)
    for ri, row in enumerate(header_rows + data_rows, start=1):
        for ci, val in enumerate(row, start=1):
            ws.cell(row=ri, column=ci, value=val)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return stats


def main() -> None:
    base = Path(__file__).parent.parent
    args = sys.argv[1:]
    base_path = Path(args[0]) if len(args) > 0 else base / "trdata" / "2026_03_Suivi_des_emplois_en_ETPT_RPROG.xlsx"
    corr_path = (
        Path(args[1])
        if len(args) > 1
        else base / "trdata" / "2026_08_Suivi_des_emplois_en_ETPT_RPROG_19082026 - Copie.xlsx"
    )
    out_path = (
        Path(args[2])
        if len(args) > 2
        else base / "trdata" / "2026_08_Suivi_des_emplois_en_ETPT_RPROG_fiabilise.xlsx"
    )

    if not base_path.exists():
        raise SystemExit(f"Fichier base introuvable : {base_path}")
    if not corr_path.exists():
        raise SystemExit(f"Fichier corrections introuvable : {corr_path}")

    stats = merge(base_path, corr_path, out_path)
    print(f"Fusion terminée → {out_path}")
    print(f"  Corrections reçues : {stats['total_corrections']}")
    print(f"  Lignes mises à jour : {stats['updated']}")
    print(f"  Lignes insérées (nouvelles) : {stats['inserted']}")
    print(f"  Déjà conformes : {stats['unchanged']}")


if __name__ == "__main__":
    main()
